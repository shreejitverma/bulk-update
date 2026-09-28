"""Read and write ``Product Listing.xlsx``.

:func:`write_template` generates the workbook the operator fills in: typed headers with hover
help, dropdowns wired to the vocabularies in :mod:`anzorlist.ingest.schema`, worked example rows,
and an instructions sheet. :func:`read_workbook` parses it back into validated
:class:`~anzorlist.ingest.row.ListingRow` objects.

Both sides are generated from :data:`~anzorlist.ingest.schema.COLUMNS`, so adding a column is a
one-line change that automatically appears in the template, the dropdowns, and the parser.

The reader is deliberately forgiving about *shape* and strict about *content*: columns may be
reordered or removed, extra columns are ignored with a warning, but a value that would produce
an Amazon rejection is an error the operator sees before anything is uploaded.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

import structlog
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from anzorlist.ingest.row import ListingRow
from anzorlist.ingest.schema import (
    COLUMNS,
    COLUMNS_BY_HEADER,
    EXAMPLE_ROWS,
    PRODUCTS_SHEET,
    README_SHEET,
    REFERENCE_SHEET,
    RESULTS_SHEET,
    Column,
    RowError,
)

log = structlog.get_logger(__name__)

# Palette. Required columns are visually distinct so an operator scanning the sheet knows
# immediately that SKU is the only cell they *must* fill.
_HDR_REQUIRED = PatternFill("solid", fgColor="1F3864")
_HDR_OPTIONAL = PatternFill("solid", fgColor="4472C4")
_HDR_NOTES = PatternFill("solid", fgColor="7F7F7F")
_GROUP_FILLS = {
    "Control": PatternFill("solid", fgColor="D9E2F3"),
    "Offer": PatternFill("solid", fgColor="E2EFD9"),
    "Copy": PatternFill("solid", fgColor="FFF2CC"),
    "Attributes": PatternFill("solid", fgColor="FBE5D6"),
    "Notes": PatternFill("solid", fgColor="EDEDED"),
}
_EXAMPLE_FONT = Font(italic=True, color="808080")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

MAX_ROWS = 5000  # dropdown validation range; generous headroom over the full Anzor catalog


@dataclass
class ReadResult:
    """Everything the reader learned, including the failures. Errors never raise mid-parse —
    the operator gets the complete list in one pass instead of fixing one cell at a time."""

    rows: list[ListingRow]
    errors: list[RowError]
    skipped: list[str]  # SKUs present but marked Include = N
    unknown_headers: list[str]

    @property
    def ok(self) -> bool:
        return not self.errors

    def included(self) -> list[ListingRow]:
        return [r for r in self.rows if r.include]


# --------------------------------------------------------------------------------------
# Writing the template
# --------------------------------------------------------------------------------------


def write_template(
    path: Path | str,
    *,
    with_examples: bool = True,
    preserve_existing: bool = True,
) -> Path:
    """Generate the fillable workbook.

    If ``preserve_existing`` and the file already has operator-entered rows, they are carried
    over into the regenerated sheet. Regenerating the template must never destroy work — the
    template can be reissued when columns change without the operator re-typing the catalog.
    """
    path = Path(path)
    carried: list[dict[str, object]] = []
    if preserve_existing and path.exists():
        carried = _read_raw_rows(path)
        if carried:
            backup = path.with_suffix(".backup.xlsx")
            backup.write_bytes(path.read_bytes())
            log.info("workbook.backup", path=str(backup), rows=len(carried))

    wb = Workbook()
    wb.remove(wb.active)  # drop the default "Sheet"

    _build_readme(wb.create_sheet(README_SHEET))
    products = wb.create_sheet(PRODUCTS_SHEET)
    _build_products(products, carried_rows=carried, with_examples=with_examples and not carried)
    _build_reference(wb.create_sheet(REFERENCE_SHEET))
    _build_results(wb.create_sheet(RESULTS_SHEET))

    wb.active = wb.index(products)  # open on the sheet they need to fill
    wb.save(path)
    log.info("workbook.written", path=str(path), carried_rows=len(carried))
    return path


def _build_products(
    ws: Worksheet, *, carried_rows: list[dict[str, object]], with_examples: bool
) -> None:
    ws.freeze_panes = "B2"  # keep SKU and the header visible while scrolling

    for idx, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        cell = ws.cell(row=1, column=idx, value=col.header)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = (
            _HDR_REQUIRED if col.required else _HDR_NOTES if col.group == "Notes" else _HDR_OPTIONAL
        )
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        if col.help:
            required_tag = "REQUIRED\n\n" if col.required else "Optional.\n\n"
            cell.comment = Comment(
                f"{col.header}\n\n{required_tag}{col.help}", "anzorlist", width=380, height=190
            )
        ws.column_dimensions[letter].width = col.width
        _attach_validation(ws, col, letter)

    ws.row_dimensions[1].height = 34

    row_no = 2
    if carried_rows:
        for raw in carried_rows:
            for idx, col in enumerate(COLUMNS, start=1):
                value = raw.get(col.header)
                if value is not None:
                    ws.cell(row=row_no, column=idx, value=value)
            row_no += 1
    elif with_examples:
        for example in EXAMPLE_ROWS:
            for idx, col in enumerate(COLUMNS, start=1):
                if col.key in example:
                    c = ws.cell(row=row_no, column=idx, value=example[col.key])  # type: ignore[arg-type]
                    c.font = _EXAMPLE_FONT
            row_no += 1

    # Tint the remaining input area by column group so the sheet reads as sections rather than
    # an undifferentiated wall of 28 columns.
    for r in range(row_no, row_no + 40):
        for idx, col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=idx)
            cell.fill = _GROUP_FILLS.get(col.group, _GROUP_FILLS["Notes"])
            cell.border = _BORDER

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"


def _attach_validation(ws: Worksheet, col: Column, letter: str) -> None:
    """Wire an Excel dropdown or numeric guard for a column, when the type admits one."""
    dv: DataValidation | None = None
    if col.choices:
        # Inline list formula. Excel caps this at 255 characters; every vocabulary here fits,
        # and the reader enforces the same set regardless, so a truncated dropdown could not
        # let a bad value through silently.
        formula = '"' + ",".join(col.choices) + '"'
        if len(formula) <= 255:
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
            dv.error = f"Choose one of: {', '.join(col.choices)}"
            dv.errorTitle = f"Invalid {col.header}"
            dv.prompt = col.help[:250]
            dv.promptTitle = col.header
    elif col.type == "int":
        dv = DataValidation(
            type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True
        )
        dv.error = "Whole number, zero or greater."
        dv.errorTitle = f"Invalid {col.header}"
    elif col.type == "decimal":
        dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
        dv.error = "A number greater than zero, or blank to use the website value."
        dv.errorTitle = f"Invalid {col.header}"

    if dv is not None:
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{MAX_ROWS}")


def _build_readme(ws: Worksheet) -> None:
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 118
    ws.sheet_view.showGridLines = False

    lines: list[tuple[str, str]] = [
        ("h1", "Anzor Jewelry → Amazon listing sheet"),
        (
            "p",
            "Fill in the Products tab. Only the SKU column is required — everything else is "
            "either optional or has a sensible default.",
        ),
        ("h2", "How it works"),
        ("li", "1.  You enter a SKU, e.g. R985."),
        (
            "li",
            "2.  The system fetches that product from anzorjewelrycorp.com and extracts the "
            "title, description, Item Details specs, price, images, and available ring sizes.",
        ),
        (
            "li",
            "3.  It writes Amazon-compliant copy from those extracted facts, prices the item "
            "with your fee markup, hosts the images, and builds the listing payload.",
        ),
        (
            "li",
            "4.  It validates that payload against Amazon's own JSON schema for the product "
            "type, offline, before anything is sent.",
        ),
        (
            "li",
            "5.  It submits in VALIDATION_PREVIEW mode first — Amazon checks the listing and "
            "returns errors without creating anything.",
        ),
        ("li", "6.  Only after you review the preview and pass --confirm does anything go live."),
        ("h2", "The only column you must fill"),
        (
            "p",
            "SKU. Everything else is an override that beats the website data. A blank cell "
            "means 'use what the site says' — it never blanks out a value.",
        ),
        ("h2", "Columns worth knowing about"),
        ("li", "Include? — set to N to keep a row in the sheet without uploading it."),
        (
            "li",
            "Marketplaces — comma-separated codes (US, CA, MX, UK, DE, ...). Blank uses your "
            ".env default. Note that the EU marketplaces need a second, separate SP-API "
            "authorization from US/CA/MX.",
        ),
        (
            "li",
            "Variations — 'site' builds a proper parent/child family from the ring sizes on "
            "the product page, so all sizes share one detail page and one review count. "
            "'none' lists a single standalone item. Use 'none' for earrings and pendants.",
        ),
        (
            "li",
            "Price Override — leave blank. The default price is the website price plus your "
            "PRICE_MARKUP_AMAZON (20%), which absorbs Amazon's referral fee. Filling this in "
            "bypasses the markup entirely and sets the exact price.",
        ),
        (
            "li",
            "UPC / EAN — leave blank. Blank routes the SKU through GTIN exemption, which is "
            "the correct path for jewelry you manufacture. Never invent a barcode: Amazon "
            "validates them against the GS1 registry and a fake one can suspend the account.",
        ),
        (
            "li",
            "Title / Bullet / Description overrides — leave blank unless you have a reason. "
            "Generated copy is checked against the extracted specs so it cannot claim a carat "
            "weight, metal purity, or stone origin the website does not state. Anything you "
            "type here bypasses that check.",
        ),
        ("h2", "Hover any column header for its full explanation."),
        ("h2", "What to run"),
        ("mono", "anzorlist doctor                    # check credentials and config"),
        ("mono", "anzorlist workbook validate         # check this sheet, no network"),
        ("mono", "anzorlist build                     # extract, price, copy, build payloads"),
        ("mono", "anzorlist amazon preflight          # category gating + product-type check"),
        ("mono", "anzorlist amazon validate           # Amazon's own dry-run, creates nothing"),
        ("mono", "anzorlist amazon submit --confirm   # the only command that creates listings"),
        ("h2", "Safety"),
        (
            "p",
            "Nothing is created on Amazon without --confirm, and nothing becomes buyable "
            "without a second explicit step. Every submission is recorded in a local ledger "
            "with the exact payload sent, so any listing can be traced or rolled back.",
        ),
        ("h2", "Results"),
        (
            "p",
            "After a run, the Upload Results tab is filled in with the status, submission ID, "
            "and any Amazon-reported issue for each SKU.",
        ),
    ]

    styles = {
        "h1": (Font(bold=True, size=18, color="1F3864"), 30),
        "h2": (Font(bold=True, size=13, color="1F3864"), 24),
        "p": (Font(size=11), None),
        "li": (Font(size=11), None),
        "mono": (Font(size=10, name="Consolas", color="333333"), None),
    }
    r = 2
    for kind, text in lines:
        font, height = styles[kind]
        cell = ws.cell(row=r, column=2, value=text)
        cell.font = font
        cell.alignment = Alignment(
            wrap_text=True, vertical="top", indent=2 if kind in ("li", "mono") else 0
        )
        if height:
            ws.row_dimensions[r].height = height
        elif len(text) > 110:
            ws.row_dimensions[r].height = 15 * (len(text) // 110 + 1)
        if kind == "mono":
            cell.fill = PatternFill("solid", fgColor="F2F2F2")
        r += 1


def _build_reference(ws: Worksheet) -> None:
    """Vocabularies, for lookup and for any operator who wants to build their own formulas."""
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 90

    ws["A1"] = "Column"
    ws["B1"] = "Allowed values"
    ws["C1"] = "Meaning"
    for c in ("A1", "B1", "C1"):
        ws[c].font = Font(bold=True, color="FFFFFF")
        ws[c].fill = _HDR_OPTIONAL

    r = 2
    for col in COLUMNS:
        if not col.choices:
            continue
        for i, choice in enumerate(col.choices):
            ws.cell(row=r, column=1, value=col.header if i == 0 else "")
            ws.cell(row=r, column=2, value=choice)
            if i == 0:
                cell = ws.cell(row=r, column=3, value=col.help)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1

    from anzorlist.marketplaces import BY_CODE

    ws.cell(row=r, column=1, value="Marketplace codes").font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value="Code").font = Font(bold=True)
    ws.cell(row=r, column=2, value="Marketplace ID").font = Font(bold=True)
    ws.cell(row=r, column=3, value="Country / region / currency").font = Font(bold=True)
    r += 1
    for code, m in sorted(BY_CODE.items()):
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=m.marketplace_id)
        ws.cell(
            row=r,
            column=3,
            value=f"{m.country} — {m.region.value.upper()} region, {m.currency}, {m.domain}",
        )
        r += 1


def _build_results(ws: Worksheet) -> None:
    headers = [
        "SKU",
        "Marketplace",
        "Child SKU",
        "Status",
        "Submission ID",
        "Issue Code",
        "Severity",
        "Message",
        "Payload",
        "Timestamp",
    ]
    widths = [14, 12, 16, 14, 30, 22, 10, 70, 44, 22]
    for i, (h, w) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HDR_OPTIONAL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.cell(
        row=2,
        column=1,
        value="(filled in automatically after `anzorlist amazon validate` or `submit`)",
    ).font = _EXAMPLE_FONT


# --------------------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------------------


def _read_raw_rows(path: Path) -> list[dict[str, object]]:
    """Header-keyed raw rows from the Products sheet. No validation; used by the regenerator."""
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # a corrupt or non-xlsx file
        log.warning("workbook.unreadable", path=str(path), error=str(exc))
        return []
    if PRODUCTS_SHEET not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[PRODUCTS_SHEET]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        wb.close()
        return []
    headers = [str(h).strip() if h is not None else "" for h in header]
    out: list[dict[str, object]] = []
    for values in rows_iter:
        record = {h: v for h, v in zip(headers, values, strict=False) if h and v is not None}
        if record.get("SKU"):
            out.append(record)
    wb.close()
    return out


def read_workbook(path: Path | str) -> ReadResult:
    """Parse and validate the Products sheet. Collects every error rather than raising on the
    first, so one pass through the sheet fixes everything."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(
            f"{path} not found. Run `anzorlist workbook init` to generate the template."
        )

    wb = load_workbook(path, data_only=True, read_only=True)
    if PRODUCTS_SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"{path} has no {PRODUCTS_SHEET!r} sheet (found: {', '.join(wb.sheetnames)}). "
            f"Run `anzorlist workbook init` to regenerate it — existing rows are preserved."
        )

    ws = wb[PRODUCTS_SHEET]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_values = next(rows_iter)
    except StopIteration:
        wb.close()
        return ReadResult([], [], [], [])

    headers = [str(h).strip() if h is not None else "" for h in header_values]
    unknown = [h for h in headers if h and h not in COLUMNS_BY_HEADER]
    index_of: dict[str, int] = {h: i for i, h in enumerate(headers) if h in COLUMNS_BY_HEADER}

    missing_required = [h for h in ("SKU",) if h not in index_of]
    if missing_required:
        wb.close()
        raise ValueError(
            f"{path}!{PRODUCTS_SHEET} is missing required column(s): "
            f"{', '.join(missing_required)}. Run `anzorlist workbook init` to regenerate."
        )

    rows: list[ListingRow] = []
    errors: list[RowError] = []
    skipped: list[str] = []
    seen_skus: dict[str, int] = {}

    for offset, values in enumerate(rows_iter):
        excel_row = offset + 2
        payload, cell_errors = _row_payload(values, index_of, excel_row)
        if payload is None:
            continue  # entirely blank row
        errors.extend(cell_errors)
        if cell_errors:
            continue

        try:
            row = ListingRow(source_row=excel_row, **payload)  # type: ignore[arg-type]
        except ValidationError as exc:
            for err in exc.errors():
                key = str(err["loc"][0]) if err["loc"] else "?"
                col = next((c for c in COLUMNS if c.key == key), None)
                header = col.header if col else key
                errors.append(
                    RowError(
                        row=excel_row,
                        header=header,
                        value=payload.get(key),
                        message=err["msg"].removeprefix("Value error, "),
                        cell=_cell_ref(index_of.get(header), excel_row),
                    )
                )
            continue

        if row.sku in seen_skus:
            errors.append(
                RowError(
                    row=excel_row,
                    header="SKU",
                    value=row.sku,
                    message=f"duplicate SKU — already listed on row {seen_skus[row.sku]}. "
                    f"Amazon keys listings on seller SKU, so two rows would overwrite "
                    f"each other.",
                    cell=_cell_ref(index_of.get("SKU"), excel_row),
                )
            )
            continue
        seen_skus[row.sku] = excel_row

        if not row.include:
            skipped.append(row.sku)
        rows.append(row)

    wb.close()
    if unknown:
        log.warning("workbook.unknown_columns", columns=unknown)
    log.info(
        "workbook.read",
        path=str(path),
        rows=len(rows),
        included=sum(1 for r in rows if r.include),
        skipped=len(skipped),
        errors=len(errors),
    )
    return ReadResult(rows=rows, errors=errors, skipped=skipped, unknown_headers=unknown)


def _cell_ref(col_index: int | None, row: int) -> str:
    return f"{get_column_letter(col_index + 1)}{row}" if col_index is not None else ""


def _row_payload(
    values: tuple[object, ...], index_of: dict[str, int], excel_row: int
) -> tuple[dict[str, object] | None, list[RowError]]:
    """Coerce one spreadsheet row into kwargs for :class:`ListingRow`."""
    payload: dict[str, object] = {}
    errors: list[RowError] = []

    sku_raw = values[index_of["SKU"]] if index_of["SKU"] < len(values) else None
    if sku_raw is None or not str(sku_raw).strip():
        return None, []  # blank row, or a trailing formatted-but-empty row

    for header, idx in index_of.items():
        col = COLUMNS_BY_HEADER[header]
        raw = values[idx] if idx < len(values) else None
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            continue
        try:
            payload[col.key] = _coerce(col, raw)
        except (ValueError, InvalidOperation, ArithmeticError) as exc:
            errors.append(
                RowError(
                    row=excel_row,
                    header=header,
                    value=raw,
                    message=str(exc),
                    cell=_cell_ref(idx, excel_row),
                )
            )
    return payload, errors


def _coerce(col: Column, raw: object) -> object:
    """Spreadsheet cell -> Python value, with messages an operator can act on."""
    if col.type == "bool":
        if isinstance(raw, bool):
            return raw
        text = str(raw).strip().lower()
        if text in ("y", "yes", "true", "1", "x", "✓"):
            return True
        if text in ("n", "no", "false", "0", ""):
            return False
        raise ValueError("expected Y or N")

    if col.type == "int":
        if isinstance(raw, bool):
            raise ValueError("expected a whole number")
        value = Decimal(str(raw).strip())
        if value != value.to_integral_value():
            raise ValueError(f"expected a whole number, not {value}")
        return int(value)

    if col.type == "decimal":
        text = str(raw).replace("$", "").replace(",", "").strip()
        try:
            return Decimal(text)
        except InvalidOperation:
            raise ValueError("expected a number, e.g. 1295.00") from None

    if col.type == "list":
        return [part.strip().upper() for part in str(raw).split(",") if part.strip()]

    text = str(raw).strip()
    # Excel stores whole numbers typed into text columns as floats; "925.0" as a metal stamp is
    # wrong in a way that is invisible in the UI, so normalise it back.
    if isinstance(raw, float) and raw.is_integer():
        text = str(int(raw))
    return text
