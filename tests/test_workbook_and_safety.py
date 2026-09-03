"""Workbook round-trip, ledger behaviour, and the live-write safety gates.

The safety tests matter most: they assert that the system *refuses* to do something. A
regression that removes a gate is silent — nothing fails, listings just go live unreviewed.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from anzorlist.ingest import read_workbook, write_template
from anzorlist.ingest.row import ListingRow
from anzorlist.ingest.schema import COLUMNS, PRODUCTS_SHEET, README_SHEET
from anzorlist.models.listing import BuiltListing, ListingStatus, OfferTerms, SubmissionOutcome
from anzorlist.store.db import Ledger


class TestWorkbookRoundTrip:
    def test_template_has_the_expected_sheets(self, tmp_path: Path):
        path = write_template(tmp_path / "wb.xlsx")
        wb = load_workbook(path)
        assert PRODUCTS_SHEET in wb.sheetnames
        assert README_SHEET in wb.sheetnames

    def test_every_column_appears_as_a_header(self, tmp_path: Path):
        path = write_template(tmp_path / "wb.xlsx")
        wb = load_workbook(path)
        headers = [c.value for c in wb[PRODUCTS_SHEET][1]]
        for column in COLUMNS:
            assert column.header in headers

    def test_example_rows_parse(self, tmp_path: Path):
        path = write_template(tmp_path / "wb.xlsx", with_examples=True)
        result = read_workbook(path)
        assert result.ok, [str(e) for e in result.errors]
        assert len(result.rows) == 4
        assert {r.sku for r in result.rows} == {"R985", "E1154", "S220", "E711"}

    def test_include_flag_partitions_the_rows(self, tmp_path: Path):
        path = write_template(tmp_path / "wb.xlsx", with_examples=True)
        result = read_workbook(path)
        assert len(result.included()) == 2
        assert set(result.skipped) == {"S220", "E711"}

    def test_regeneration_preserves_operator_rows(self, tmp_path: Path):
        """Reissuing the template when columns change must never lose typed data."""
        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        wb = load_workbook(path)
        ws = wb[PRODUCTS_SHEET]
        ws.cell(row=2, column=1, value="R9001")
        ws.cell(row=3, column=1, value="E9002")
        wb.save(path)

        write_template(path, preserve_existing=True)
        result = read_workbook(path)
        assert {r.sku for r in result.rows} == {"R9001", "E9002"}

    def test_missing_file_says_what_to_run(self, tmp_path: Path):
        with pytest.raises(FileNotFoundError, match="workbook init"):
            read_workbook(tmp_path / "nope.xlsx")

    def test_duplicate_skus_are_rejected(self, tmp_path: Path):
        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        wb = load_workbook(path)
        ws = wb[PRODUCTS_SHEET]
        ws.cell(row=2, column=1, value="R985")
        ws.cell(row=3, column=1, value="R985")
        wb.save(path)
        result = read_workbook(path)
        assert any("duplicate" in e.message for e in result.errors)

    def test_errors_carry_a_cell_reference(self, tmp_path: Path):
        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        wb = load_workbook(path)
        ws = wb[PRODUCTS_SHEET]
        ws.cell(row=2, column=1, value="R985")
        condition_col = next(i for i, c in enumerate(COLUMNS, 1) if c.key == "condition")
        ws.cell(row=2, column=condition_col, value="brand_spanking_new")
        wb.save(path)
        result = read_workbook(path)
        assert result.errors
        assert result.errors[0].cell, "an error must point at a cell the operator can find"


class TestListingRow:
    def test_sku_is_uppercased(self):
        assert ListingRow(sku=" r985 ").sku == "R985"

    def test_unknown_condition_is_rejected(self):
        with pytest.raises(ValueError, match="unknown condition"):
            ListingRow(sku="R985", condition="mint")

    def test_unknown_product_type_is_rejected_early(self):
        with pytest.raises(ValueError, match="not in the known jewelry product-type list"):
            ListingRow(sku="R985", amazon_product_type="WIDGET")

    def test_negative_quantity_is_rejected(self):
        with pytest.raises(ValueError):
            ListingRow(sku="R985", quantity=-1)

    @pytest.mark.parametrize("gtin", ["036000291452", "4006381333931"])
    def test_valid_gtins_pass(self, gtin: str):
        assert ListingRow(sku="R985", upc_ean=gtin).upc_ean == gtin

    @pytest.mark.parametrize("bad", ["12345", "abcdefghijkl", "036000291453"])
    def test_invalid_gtins_are_rejected(self, bad: str):
        with pytest.raises(ValueError):
            ListingRow(sku="R985", upc_ean=bad)

    def test_blank_gtin_means_exemption(self):
        assert ListingRow(sku="R985").uses_gtin_exemption is True


def _listing(sku: str = "R985", payload_hash: str = "abc123") -> BuiltListing:
    return BuiltListing(
        sku=sku, source_sku=sku,
        marketplace_id="ATVPDKIKX0DER", marketplace_code="US",
        product_type="RING", attributes={"brand": [{"value": "Anzor"}]},
        offer=OfferTerms(price=Decimal("1250.00")),
        payload_hash=payload_hash, content_hash="0" * 64,
    )


class TestLedger:
    def test_records_and_reads_back_a_listing(self, tmp_path: Path):
        with Ledger(tmp_path / "l.sqlite") as ledger:
            ledger.record_listing(_listing())
            row = ledger.get_listing("R985", "ATVPDKIKX0DER")
            assert row is not None
            assert row["product_type"] == "RING"

    def test_unsubmitted_listing_needs_submission(self, tmp_path: Path):
        with Ledger(tmp_path / "l.sqlite") as ledger:
            assert ledger.needs_submission(_listing()) is True

    def test_identical_payload_is_not_resubmitted(self, tmp_path: Path):
        """Idempotency: an unchanged payload must not burn SP-API quota."""
        with Ledger(tmp_path / "l.sqlite") as ledger:
            listing = _listing()
            ledger.record_listing(listing)
            ledger.start_run("r1", "submit", "SUBMIT")
            ledger.record_submission(SubmissionOutcome(
                sku=listing.sku, marketplace_id=listing.marketplace_id, marketplace_code="US",
                mode="SUBMIT", status=ListingStatus.ACCEPTED, payload_hash=listing.payload_hash,
            ), "r1")
            assert ledger.needs_submission(listing) is False

    def test_changed_payload_is_resubmitted(self, tmp_path: Path):
        with Ledger(tmp_path / "l.sqlite") as ledger:
            first = _listing(payload_hash="hash-v1")
            ledger.record_listing(first)
            ledger.start_run("r1", "submit", "SUBMIT")
            ledger.record_submission(SubmissionOutcome(
                sku=first.sku, marketplace_id=first.marketplace_id, marketplace_code="US",
                mode="SUBMIT", status=ListingStatus.ACCEPTED, payload_hash="hash-v1",
            ), "r1")
            assert ledger.needs_submission(_listing(payload_hash="hash-v2")) is True

    def test_live_skus_is_the_rollback_list(self, tmp_path: Path):
        with Ledger(tmp_path / "l.sqlite") as ledger:
            listing = _listing()
            ledger.record_listing(listing)
            ledger.start_run("r1", "submit", "SUBMIT")
            ledger.record_submission(SubmissionOutcome(
                sku=listing.sku, marketplace_id=listing.marketplace_id, marketplace_code="US",
                mode="SUBMIT", status=ListingStatus.SUBMITTED, payload_hash=listing.payload_hash,
            ), "r1")
            assert [e.sku for e in ledger.live_skus()] == ["R985"]

    def test_validation_preview_does_not_mark_a_listing_live(self, tmp_path: Path):
        """A dry run must never make a SKU look live in the ledger."""
        with Ledger(tmp_path / "l.sqlite") as ledger:
            listing = _listing()
            ledger.record_listing(listing)
            ledger.start_run("r1", "validate", "VALIDATION_PREVIEW")
            ledger.record_submission(SubmissionOutcome(
                sku=listing.sku, marketplace_id=listing.marketplace_id, marketplace_code="US",
                mode="VALIDATION_PREVIEW", status=ListingStatus.VALIDATED,
                payload_hash=listing.payload_hash,
            ), "r1")
            assert ledger.live_skus() == []
            assert ledger.needs_submission(listing) is True


class TestLiveWriteGates:
    """Both gates must be open. Neither alone may cause a write."""

    def _client(self, settings):
        from unittest.mock import MagicMock

        from anzorlist.channels.amazon.listings import ListingsClient

        transport = MagicMock()
        transport.seller_id = "A1SELLER"
        return ListingsClient(transport, settings), transport

    def test_submit_without_confirm_is_refused(self, settings, us):
        from anzorlist.channels.amazon.listings import LiveWriteBlocked

        settings.allow_live = True
        client, transport = self._client(settings)
        with pytest.raises(LiveWriteBlocked, match="confirm=True"):
            client.put(_listing(), us, mode="SUBMIT", confirm=False)
        transport.request.assert_not_called()

    def test_submit_without_allow_live_is_refused(self, settings, us):
        from anzorlist.channels.amazon.listings import LiveWriteBlocked

        settings.allow_live = False
        client, transport = self._client(settings)
        with pytest.raises(LiveWriteBlocked, match="ANZOR_ALLOW_LIVE"):
            client.put(_listing(), us, mode="SUBMIT", confirm=True)
        transport.request.assert_not_called()

    def test_validation_preview_needs_no_gate(self, settings, us):
        settings.allow_live = False
        client, transport = self._client(settings)
        transport.request.return_value.json = {"sku": "R985", "status": "ACCEPTED", "issues": []}
        transport.request.return_value.request_id = "req-1"
        transport.request.return_value.status = 200
        outcome = client.put(_listing(), us, mode="VALIDATION_PREVIEW")
        assert outcome.status is ListingStatus.VALIDATED
        assert transport.request.call_args.kwargs["params"]["mode"] == "VALIDATION_PREVIEW"

    def test_submit_mode_omits_the_preview_parameter(self, settings, us):
        settings.allow_live = True
        client, transport = self._client(settings)
        transport.request.return_value.json = {"sku": "R985", "status": "ACCEPTED", "issues": []}
        transport.request.return_value.request_id = "req-2"
        transport.request.return_value.status = 200
        client.put(_listing(), us, mode="SUBMIT", confirm=True)
        assert "mode" not in transport.request.call_args.kwargs["params"]

    def test_delete_requires_both_gates(self, settings, us):
        from anzorlist.channels.amazon.listings import LiveWriteBlocked

        settings.allow_live = True
        client, transport = self._client(settings)
        with pytest.raises(LiveWriteBlocked):
            client.delete("R985", us, confirm=False)
        transport.request.assert_not_called()

    def test_amazon_reported_errors_become_a_failed_outcome(self, settings, us):
        client, transport = self._client(settings)
        transport.request.return_value.json = {
            "sku": "R985", "status": "INVALID",
            "issues": [{"code": "90220", "message": "brand is required",
                        "severity": "ERROR", "attributeNames": ["brand"]}],
        }
        transport.request.return_value.request_id = "req-3"
        transport.request.return_value.status = 200
        outcome = client.put(_listing(), us, mode="VALIDATION_PREVIEW")
        assert outcome.status is ListingStatus.VALIDATION_FAILED
        assert not outcome.accepted
        assert outcome.issues[0].attribute_names == ["brand"]

    def test_unknown_severity_is_treated_as_blocking(self, settings, us):
        """Fail closed: an unrecognised severity must never be silently ignored."""
        client, transport = self._client(settings)
        transport.request.return_value.json = {
            "sku": "R985", "status": "ACCEPTED",
            "issues": [{"code": "X", "message": "?", "severity": "MYSTERY"}],
        }
        transport.request.return_value.request_id = "req-4"
        transport.request.return_value.status = 200
        outcome = client.put(_listing(), us, mode="VALIDATION_PREVIEW")
        assert outcome.issues[0].blocking is True
