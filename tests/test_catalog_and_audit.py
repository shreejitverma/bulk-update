"""Catalog enumeration, image auditing, and the workbook bulk-population helpers.

All offline. The scanner and auditor are exercised against a fake transport, because the
behaviour worth pinning down is how they cope with a *misbehaving* server — a paginator that
clamps out-of-range pages, cross-listed duplicates, missing images — and a live site cannot be
made to misbehave on demand.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from anzorlist.extract.catalog import (
    EMPTY_PAGE_TOLERANCE,
    TOP_LEVEL_CATEGORIES,
    CatalogScanner,
)


class FakeResponse:
    def __init__(self, text: str) -> None:
        self.content = text.encode("cp1252", "replace")
        self.status_code = 200


class FakeSite:
    """Stands in for SiteClient. Serves canned category pages and counts requests."""

    base_url = "https://example.test"

    def __init__(self, pages: dict[str, str]) -> None:
        self.pages = pages
        self.requested: list[str] = []

    def _get(self, url: str) -> FakeResponse:
        self.requested.append(url)
        if url not in self.pages:
            raise KeyError(f"no canned page for {url}")
        return FakeResponse(self.pages[url])

    @staticmethod
    def _decode(raw: bytes) -> tuple[str, str]:
        return raw.decode("cp1252", "replace"), "cp1252"


def page(skus: list[str], total_pages: int | None = None) -> str:
    links = " ".join(f'<a href="prodview.asp?SKU={s}">{s}</a>' for s in skus)
    counter = f"of {total_pages}&nbsp;" if total_pages else ""
    return f"<html><body>{links} {counter}</body></html>"


def url(cid: int, p: int) -> str:
    base = f"https://example.test/Scripts/prodList.asp?idCategory={cid}"
    return base if p <= 1 else f"{base}&curPage={p}"


class TestCatalogScanner:
    def test_walks_every_page_the_site_reports(self):
        site = FakeSite(
            {
                url(15, 1): page(["R1", "R2"], total_pages=3),
                url(15, 2): page(["R3", "R4"]),
                url(15, 3): page(["R5"]),
            }
        )
        scan = CatalogScanner(site).scan_category(15)
        assert scan.skus == ["R1", "R2", "R3", "R4", "R5"]
        assert scan.pages_fetched == 3
        assert scan.reported_pages == 3
        assert scan.complete and not scan.truncated

    def test_stops_when_the_paginator_clamps_instead_of_erroring(self):
        """Classic ASP often returns the last valid page forever for an out-of-range curPage.

        Terminating on "no new SKUs" rather than on a 404 is what makes the walk finite here.
        """
        last = page(["R9", "R10"])
        site = FakeSite({url(15, 1): page(["R1"]), **{url(15, p): last for p in range(2, 40)}})
        scan = CatalogScanner(site).scan_category(15)
        assert scan.skus == ["R1", "R9", "R10"]
        assert scan.pages_fetched <= 2 + EMPTY_PAGE_TOLERANCE
        assert not scan.truncated

    def test_tolerates_one_empty_page_without_giving_up(self):
        site = FakeSite(
            {
                url(15, 1): page(["R1"]),
                url(15, 2): page([]),
                url(15, 3): page(["R2"]),
                url(15, 4): page([]),
                url(15, 5): page([]),
                url(15, 6): page([]),
            }
        )
        scan = CatalogScanner(site).scan_category(15)
        assert scan.skus == ["R1", "R2"], "a single gap must not end the walk"

    def test_partial_results_survive_a_mid_walk_failure(self):
        site = FakeSite({url(15, 1): page(["R1", "R2"], total_pages=5)})  # page 2 raises
        scan = CatalogScanner(site).scan_category(15)
        assert scan.skus == ["R1", "R2"]
        assert scan.truncated
        assert not scan.complete

    def test_skus_are_uppercased_and_deduped_within_a_page(self):
        site = FakeSite(
            {
                url(15, 1): page(["r1", "R1", "r2"]),
                url(15, 2): page([]),
                url(15, 3): page([]),
                url(15, 4): page([]),
            }
        )
        assert CatalogScanner(site).scan_category(15).skus == ["R1", "R2"]

    def test_cross_listed_skus_are_deduped_across_categories(self):
        """A product legitimately appears in several departments; Amazon keys on seller SKU,
        so shipping it twice would be a collision."""
        empty = page([])
        site = FakeSite(
            {
                url(15, 1): page(["R1", "SHARED"]),
                url(15, 2): empty,
                url(15, 3): empty,
                url(15, 4): empty,
                url(16, 1): page(["E1", "SHARED"]),
                url(16, 2): empty,
                url(16, 3): empty,
                url(16, 4): empty,
            }
        )
        scan = CatalogScanner(site).scan([15, 16])
        assert scan.skus == ["R1", "SHARED", "E1"]
        assert scan.category_of()["SHARED"] == TOP_LEVEL_CATEGORIES[15]

    def test_by_prefix_counts_families(self):
        empty = page([])
        site = FakeSite(
            {
                url(15, 1): page(["R1", "R2", "E1"]),
                url(15, 2): empty,
                url(15, 3): empty,
                url(15, 4): empty,
            }
        )
        assert CatalogScanner(site).scan([15]).by_prefix() == {"E": 1, "R": 2}


class TestImageAudit:
    def _report(self, sizes: dict[str, tuple[int, int]]):
        """Build an AuditReport directly — the HTTP probing is not what's under test here."""
        from anzorlist.media.audit import AuditReport, ImageFact, SkuAudit

        report = AuditReport()
        for sku, slots in sizes.items():  # type: ignore[assignment]
            audit = SkuAudit(sku=sku)
            for slot, (w, h) in slots.items():  # type: ignore[attr-defined]
                audit.images.append(
                    ImageFact(
                        sku=sku,
                        slot=slot,
                        url=f"/{sku}{slot}.jpg",
                        width=w,
                        height=h,
                        status=200,
                        image_format="JPEG",
                    )
                )
            report.skus.append(audit)
        return report

    def test_main_image_below_1000px_blocks_the_sku(self):
        report = self._report({"R1": {"a": (400, 400), "b": (400, 400)}})
        assert report.skus[0].status == "main_too_small"
        assert not report.skus[0].listable
        assert report.summary()["skus_blocked"] == 1

    def test_compliant_main_makes_the_sku_listable_even_with_small_alternates(self):
        """Amazon requires the main image to clear the bar; alternates are optional."""
        report = self._report({"E1": {"a": (1000, 1000), "b": (400, 400)}})
        assert report.skus[0].listable
        assert report.summary()["skus_listable"] == 1

    def test_missing_main_image_is_reported_distinctly(self):
        from anzorlist.media.audit import AuditReport, SkuAudit

        report = AuditReport(skus=[SkuAudit(sku="X1")])
        assert report.skus[0].status == "no_main_image"
        assert not report.skus[0].listable

    def test_verdict_buckets(self):
        report = self._report(
            {"A": {"a": (1600, 1600)}, "B": {"a": (1000, 1000)}, "C": {"a": (400, 400)}}
        )
        summary = report.summary()
        assert summary["images_ready"] == 1
        assert summary["images_at_minimum"] == 1
        assert summary["images_too_small"] == 1

    def test_histogram_shows_where_images_sit(self):
        report = self._report(
            {"A": {"a": (400, 400)}, "B": {"a": (400, 400)}, "C": {"a": (1600, 1600)}}
        )
        assert report.size_histogram() == {"<500px": 2, "1600px+": 1}

    def test_csv_is_a_usable_worklist(self, tmp_path: Path):
        import csv

        report = self._report({"R1": {"a": (400, 400)}, "E1": {"a": (1200, 1200)}})
        out = report.write_csv(tmp_path / "audit.csv")
        rows = list(csv.DictReader(out.open()))
        by_sku = {r["sku"]: r for r in rows}
        assert by_sku["R1"]["needs_new_photography"] == "yes"
        assert by_sku["E1"]["needs_new_photography"] == "no"


class TestWorkbookBulkPopulation:
    def test_appends_new_skus_staged_off_by_default(self, tmp_path: Path):
        """Discovery must not be the same act as queueing thousands of uploads."""
        from anzorlist.ingest import read_workbook, write_template
        from anzorlist.ingest.workbook import append_skus

        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        added, skipped = append_skus(path, ["R100", "R101", "E200"])
        assert (added, skipped) == (3, 0)

        result = read_workbook(path)
        assert {r.sku for r in result.rows} == {"R100", "R101", "E200"}
        assert result.included() == [], "appended rows must default to Include = N"

    def test_rerunning_a_scan_only_adds_what_is_new(self, tmp_path: Path):
        from anzorlist.ingest import write_template
        from anzorlist.ingest.workbook import append_skus

        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        append_skus(path, ["R100", "R101"])
        added, skipped = append_skus(path, ["R100", "R101", "R102"])
        assert (added, skipped) == (1, 2)

    def test_existing_operator_edits_are_never_touched(self, tmp_path: Path):
        from openpyxl import load_workbook

        from anzorlist.ingest import read_workbook, write_template
        from anzorlist.ingest.workbook import append_skus

        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        append_skus(path, ["R100"], include=True)
        wb = load_workbook(path)
        ws = wb["Products"]
        headers = [c.value for c in ws[1]]
        ws.cell(row=2, column=headers.index("Price Override (USD)") + 1, value=1234.56)
        wb.save(path)

        append_skus(path, ["R100", "R200"])
        rows = {r.sku: r for r in read_workbook(path).rows}
        assert str(rows["R100"].price_override_usd) == "1234.56"
        assert rows["R100"].include is True, "an existing row's Include must not be reset"

    def test_include_flag_opts_rows_in(self, tmp_path: Path):
        from anzorlist.ingest import read_workbook, write_template
        from anzorlist.ingest.workbook import append_skus

        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        append_skus(path, ["R100"], include=True)
        assert len(read_workbook(path).included()) == 1

    def test_results_are_written_back_to_the_sheet(self, tmp_path: Path):
        from openpyxl import load_workbook

        from anzorlist.ingest import write_template
        from anzorlist.ingest.workbook import write_results
        from anzorlist.models.listing import ListingIssue, ListingStatus, SubmissionOutcome

        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        outcomes = [
            SubmissionOutcome(
                sku="R100",
                marketplace_id="ATVPDKIKX0DER",
                marketplace_code="US",
                mode="VALIDATION_PREVIEW",
                status=ListingStatus.VALIDATED,
            ),
            SubmissionOutcome(
                sku="R101",
                marketplace_id="ATVPDKIKX0DER",
                marketplace_code="US",
                mode="VALIDATION_PREVIEW",
                status=ListingStatus.VALIDATION_FAILED,
                issues=[ListingIssue(code="90220", message="brand is required", source="amazon")],
            ),
        ]
        assert write_results(path, outcomes) == 2

        ws = load_workbook(path)["Upload Results"]
        assert ws.cell(row=2, column=1).value == "R100"
        assert ws.cell(row=3, column=6).value == "90220"
        assert "brand is required" in str(ws.cell(row=3, column=8).value)

    def test_results_sheet_is_replaced_not_appended(self, tmp_path: Path):
        from openpyxl import load_workbook

        from anzorlist.ingest import write_template
        from anzorlist.ingest.workbook import write_results
        from anzorlist.models.listing import ListingStatus, SubmissionOutcome

        path = write_template(tmp_path / "wb.xlsx", with_examples=False)
        make = lambda sku: SubmissionOutcome(  # noqa: E731
            sku=sku,
            marketplace_id="ATVPDKIKX0DER",
            marketplace_code="US",
            mode="SUBMIT",
            status=ListingStatus.SUBMITTED,
        )
        write_results(path, [make("A"), make("B"), make("C")])
        write_results(path, [make("Z")])
        ws = load_workbook(path)["Upload Results"]
        assert ws.cell(row=2, column=1).value == "Z"
        assert ws.cell(row=3, column=1).value in (None, "")


@pytest.mark.parametrize("cid,name", sorted(TOP_LEVEL_CATEGORIES.items()))
def test_top_level_categories_are_named(cid: int, name: str):
    assert isinstance(cid, int) and name and not name.isdigit()
