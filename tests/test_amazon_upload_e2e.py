"""End to end: workbook -> build -> amazon validate / submit / feed-status, against a fake Amazon.

These drive the real CLI. The only substitutions are at the process boundary: the SP-API client
pool gets a transport that routes to :class:`FakeAmazon`, the Anzor site is served from the
committed HTML fixtures through the page cache, and image hosting is a stub uploader. Everything
between - artifacts on disk, the ledger, planning, gating, feed reconciliation - is the code that
runs in production.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
from pathlib import Path
from typing import Any

import pytest
import structlog
from conftest import FIXTURES
from fake_amazon import FakeAmazon
from openpyxl import load_workbook
from PIL import Image
from typer.testing import CliRunner

from anzorlist import cli
from anzorlist.channels.amazon import client as spapi_client
from anzorlist.channels.amazon.client import ClientPool
from anzorlist.config import Settings
from anzorlist.config import settings as get_settings
from anzorlist.ingest import write_template
from anzorlist.media.pipeline import MediaPipeline, ProcessedImage
from anzorlist.store.db import Ledger

BUILT_SKUS = ("R985", "E1154")  # the example rows with Include = Y
R985_FAMILY = 30  # one parent + 29 ring sizes


class _StubUploader:
    def upload(self, image: ProcessedImage) -> str:
        return f"https://img.test/{image.content_hash}.jpg"


def _white_jpeg(size: int = 1200) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (size, size), (255, 255, 255)).save(buf, format="JPEG")
    return buf.getvalue()


@pytest.fixture
def env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    """An isolated operator workspace with credentials, a workbook, cached pages, and photos."""
    monkeypatch.chdir(tmp_path)
    data = tmp_path / "data"
    for key, value in {
        "ANZOR_DATA_DIR": str(data),
        "ANZOR_STATE_DB": str(data / "state.sqlite"),
        "ANZOR_WORKBOOK": str(tmp_path / "Product Listing.xlsx"),
        "ANZOR_IMAGES_DIR": str(tmp_path / "images"),
        "ANZOR_MARKETPLACES": "US",
        "SPAPI_LWA_CLIENT_ID": "amzn1.application-oa2-client.test",
        "SPAPI_LWA_CLIENT_SECRET": "secret",
        "SPAPI_REFRESH_TOKEN_NA": "Atzr|test",
        "SPAPI_SELLER_ID_NA": "A1SELLER",
        "ANZOR_ALLOW_LIVE": "true",
    }.items():
        monkeypatch.setenv(key, value)
    get_settings.cache_clear()

    write_template(tmp_path / "Product Listing.xlsx", with_examples=True)
    for sku in BUILT_SKUS:
        raw = (FIXTURES / f"{sku}.html").read_bytes()
        cache = data / "raw" / sku
        cache.mkdir(parents=True)
        (cache / "prodview.html").write_bytes(raw)
        (cache / "prodview.sha256").write_text(hashlib.sha256(raw).hexdigest())
        photos = tmp_path / "images" / sku
        photos.mkdir(parents=True)
        (photos / "01.jpg").write_bytes(_white_jpeg())

    fake = FakeAmazon()
    monkeypatch.setattr(cli, "_make_pool", lambda s: ClientPool(s, transport=fake.transport))
    monkeypatch.setattr(MediaPipeline, "_get_uploader", lambda self: _StubUploader())
    # Amazon's real rate limits would make a 30-listing family take seconds; the limiter itself
    # is not under test here.
    monkeypatch.setattr(spapi_client, "DEFAULT_RATE_LIMITS", {})
    monkeypatch.setattr(spapi_client, "_FALLBACK_LIMIT", (10_000.0, 10_000))
    yield {"root": tmp_path, "data": data, "fake": fake}
    get_settings.cache_clear()


def run(*args: str, input: str | None = None) -> Any:
    try:
        return CliRunner().invoke(cli.app, list(args), input=input, catch_exceptions=False)
    finally:
        # The CLI points structlog at the runner's stderr, which is closed afterwards.
        structlog.reset_defaults()


def build_ok() -> None:
    result = run("build", "--no-copy")
    assert result.exit_code == 0, result.output


def settings() -> Settings:
    return get_settings()


def ledger_status(sku: str) -> str | None:
    with Ledger(settings().state_db) as ledger:
        row = ledger.get_listing(sku, "ATVPDKIKX0DER")
    return None if row is None else str(row["status"])


# ------------------------------------------------------------------------------ build


def test_build_writes_full_artifacts_per_source_sku(env: dict[str, Any]) -> None:
    build_ok()
    family = sorted((env["data"] / "build" / "US" / "R985").glob("*.json"))
    assert len(family) == R985_FAMILY
    parent = json.loads((env["data"] / "build" / "US" / "R985" / "R985-PARENT.json").read_text())
    assert parent["artifactVersion"] == 2
    assert parent["source_sku"] == "R985"
    assert parent["is_parent"] is True
    child = json.loads((env["data"] / "build" / "US" / "R985" / "R985-7.25.json").read_text())
    # The offer survives the round trip; the old loader dropped it.
    assert child["offer"]["price"]
    assert child["attributes"]["main_product_image_locator"][0]["media_location"].startswith(
        "https://img.test/"
    )


def test_rebuild_removes_stale_artifacts(env: dict[str, Any]) -> None:
    stale = env["data"] / "build" / "US" / "R985" / "R985-99.json"
    stale.parent.mkdir(parents=True)
    stale.write_text("{}")
    build_ok()
    assert not stale.exists()


def test_missing_images_explain_the_fix(
    env: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    shutil.rmtree(env["root"] / "images" / "R985")
    # An operator's images folder can sit deep in their home directory; the message names it.
    images = env["root"] / ("operator-workspace-" * 8) / "images"
    monkeypatch.setenv("ANZOR_IMAGES_DIR", str(images))
    get_settings.cache_clear()
    # The site's own photo is 400px; serve that size for the website image download.
    small = io.BytesIO()
    Image.new("RGB", (400, 400), (255, 255, 255)).save(small, format="JPEG")

    class _Resp:
        content = small.getvalue()

        def raise_for_status(self) -> None:
            return None

    class _Http:
        def get(self, url: str) -> _Resp:
            return _Resp()

        def close(self) -> None:
            return None

    original_init = MediaPipeline.__init__

    def init(self: MediaPipeline, s: Settings, **kw: Any) -> None:
        original_init(self, s, **kw)
        self._http = _Http()  # type: ignore[assignment]

    monkeypatch.setattr(MediaPipeline, "__init__", init)
    # Wide enough that the report prints each message on one line, so it can be read whole.
    monkeypatch.setattr(cli.console, "_width", 1000)
    result = run("build", "--no-copy", "R985")
    assert result.exit_code == 1
    # The build report (stdout; the log lines go to stderr) shows the whole fix, including the
    # folder the photos go in.
    assert str(images / "R985") in result.stdout
    parent = json.loads((env["data"] / "build" / "US" / "R985" / "R985-7.json").read_text())
    [issue] = [i for i in parent["issues"] if i["code"] == "NoMainImage"]
    assert "400px" in issue["message"]
    assert str(images / "R985") in issue["message"]
    # With every child blocked, the parent is blocked too rather than reported as ready.
    head = json.loads((env["data"] / "build" / "US" / "R985" / "R985-PARENT.json").read_text())
    assert "FamilyBlocked" in {i["code"] for i in head["issues"]}


# ------------------------------------------------------------------------------ validate


def test_validate_by_website_sku_previews_the_whole_family(env: dict[str, Any]) -> None:
    build_ok()
    result = run("amazon", "validate", "R985")
    assert result.exit_code == 0, result.output
    fake: FakeAmazon = env["fake"]
    assert len(fake.previews()) == R985_FAMILY
    assert {c.sku for c in fake.previews()} >= {"R985-PARENT", "R985-7", "R985-12"}
    assert fake.writes() == []
    assert fake.previews()[0].sku == "R985-PARENT"


# ------------------------------------------------------------------------------ submit


def test_submit_is_gated(env: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
    build_ok()
    assert run("amazon", "submit").exit_code == 2  # no --confirm
    monkeypatch.setenv("ANZOR_ALLOW_LIVE", "false")
    get_settings.cache_clear()
    assert run("amazon", "submit", "--confirm").exit_code == 2
    assert env["fake"].listing_calls == []


def test_submit_declined_at_prompt_sends_nothing(env: dict[str, Any]) -> None:
    build_ok()
    result = run("amazon", "submit", "--confirm", input="n\n")
    assert result.exit_code == 0
    assert env["fake"].listing_calls == []


def test_submit_per_item_parent_first_previewed_and_idempotent(env: dict[str, Any]) -> None:
    build_ok()
    result = run("amazon", "submit", "--confirm", input="y\n")
    assert result.exit_code == 0, result.output
    fake: FakeAmazon = env["fake"]
    writes = [c.sku for c in fake.writes()]
    assert len(writes) == R985_FAMILY + 1  # the ring family plus the E1154 standalone
    assert writes.index("R985-PARENT") < writes.index("R985-7")
    # Every write was immediately preceded by a preview of the same SKU.
    calls = fake.listing_calls
    for i, call in enumerate(calls):
        if call.mode == "SUBMIT":
            assert calls[i - 1].sku == call.sku and calls[i - 1].mode == "VALIDATION_PREVIEW"
    assert ledger_status("R985-7") == "submitted"
    assert "not yet buyable" in result.output

    again = run("amazon", "submit", "--confirm", input="y\n")
    assert again.exit_code == 0
    assert "up to date" in again.output
    assert len(fake.writes()) == len(writes)


def test_submit_skips_listings_blocked_locally(env: dict[str, Any]) -> None:
    result = run("build", "--no-copy", "--no-upload")  # checked, but no hosted image URLs
    assert result.exit_code == 1
    result = run("amazon", "submit", "--confirm", input="y\n")
    assert result.exit_code == 1
    assert env["fake"].listing_calls == []
    assert "NoMainImage" in result.output


def test_children_of_a_rejected_parent_are_not_sent(env: dict[str, Any]) -> None:
    build_ok()
    env["fake"].reject_submit.add("R985-PARENT")
    result = run("amazon", "submit", "--confirm", "R985", input="y\n")
    assert result.exit_code == 1
    assert [c.sku for c in env["fake"].writes()] == ["R985-PARENT"]
    assert "ParentNotCreated" in result.output
    assert "not yet buyable" not in result.output  # nothing was accepted


def test_rows_set_to_include_n_are_not_submitted(env: dict[str, Any]) -> None:
    build_ok()
    wb = load_workbook(env["root"] / "Product Listing.xlsx")
    ws = wb["Products"]
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[headers.index("SKU")].value == "E1154":
            row[headers.index("Include?")].value = "N"
    wb.save(env["root"] / "Product Listing.xlsx")

    result = run("amazon", "submit", "--confirm", input="y\n")
    assert result.exit_code == 0, result.output
    assert "E1154" not in {c.sku for c in env["fake"].listing_calls}


def test_an_edited_artifact_is_resubmitted_alone(env: dict[str, Any]) -> None:
    build_ok()
    assert run("amazon", "submit", "--confirm", input="y\n").exit_code == 0
    before = len(env["fake"].writes())

    path = env["data"] / "build" / "US" / "E1154" / "E1154.json"
    artifact = json.loads(path.read_text())
    artifact["attributes"]["item_name"][0]["value"] = "Edited After Review"
    path.write_text(json.dumps(artifact))

    result = run("amazon", "submit", "--confirm", input="y\n")
    assert result.exit_code == 0
    # The one-row outcome table is narrow; its title still shows the run id on one line.
    assert re.search(r"Submission submit-\d{8}T\d{12}Z-[0-9a-f]{6}\b", result.output)
    new_writes = env["fake"].writes()[before:]
    assert [c.sku for c in new_writes] == ["E1154"]
    assert new_writes[0].body["attributes"]["item_name"][0]["value"] == "Edited After Review"


# ------------------------------------------------------------------------------ feed


def test_feed_submission_reconciles_results_by_message_id(env: dict[str, Any]) -> None:
    build_ok()
    fake: FakeAmazon = env["fake"]
    fake.reject_in_feed.add("R985-8")
    result = run("amazon", "submit", "--confirm", "--feed", input="y\n")
    assert result.exit_code == 1, result.output  # one message was rejected

    # Parents go through Listings Items; everything else rides in the feed.
    assert [c.sku for c in fake.writes()] == ["R985-PARENT"]
    [feed] = fake.feeds.values()
    assert feed["request"]["marketplaceIds"] == ["ATVPDKIKX0DER"]
    fed = {m["sku"] for m in feed["document"]["messages"]}
    assert "R985-PARENT" not in fed and {"R985-7", "R985-8", "E1154"} <= fed
    assert len(fed) == R985_FAMILY  # 29 children + E1154

    assert ledger_status("R985-8") == "rejected"
    assert ledger_status("R985-7") == "submitted"
    assert ledger_status("E1154") == "submitted"


def test_feed_preview_failure_stops_the_feed(env: dict[str, Any]) -> None:
    build_ok()
    fake: FakeAmazon = env["fake"]
    fake.reject_preview.add("E1154")
    result = run("amazon", "submit", "--confirm", "--feed", input="y\n")
    assert result.exit_code == 1
    assert fake.feeds == {}
    assert fake.writes() == []  # the parent would be an empty detail page without its children
    assert "FeedPreviewFailed" in result.output


def test_unfinished_feed_is_in_flight_until_feed_status(env: dict[str, Any]) -> None:
    build_ok()
    fake: FakeAmazon = env["fake"]
    result = run("amazon", "submit", "--confirm", "--feed", "--no-wait", "E1154", input="y\n")
    assert result.exit_code == 3, result.output  # sent, result not known yet
    [feed_id] = fake.feeds
    assert ledger_status("E1154") == "pending"

    # A second submit must not send the same payload again while the feed is unreconciled.
    again = run("amazon", "submit", "--confirm", "--feed", "E1154", input="y\n")
    assert len(fake.feeds) == 1
    assert "not been reconciled" in again.output
    assert again.exit_code == 3  # not "up to date": the result is still unknown
    listed = run("amazon", "feed-status")
    assert feed_id in listed.output

    fake.feed_status = "IN_PROGRESS"
    assert run("amazon", "feed-status", feed_id).exit_code == 3
    fake.feed_status = "DONE"
    status = run("amazon", "feed-status", feed_id)
    assert status.exit_code == 0, status.output
    assert ledger_status("E1154") == "submitted"


def test_a_failed_feed_fails_every_listing_in_it(env: dict[str, Any]) -> None:
    build_ok()
    env["fake"].feed_status = "FATAL"
    result = run("amazon", "submit", "--confirm", "--feed", "E1154", input="y\n")
    assert result.exit_code == 1
    assert ledger_status("E1154") == "rejected"
    assert "FeedFatal" in result.output


def test_a_dropped_marketplace_is_not_resurrected_by_naming_the_sku(env: dict[str, Any]) -> None:
    build_ok()
    # Simulate an earlier build for a marketplace the row no longer names.
    old = env["data"] / "build" / "CA" / "E1154"
    shutil.copytree(env["data"] / "build" / "US" / "E1154", old)
    build_ok()
    assert not old.exists()


def test_a_child_named_alone_is_held_when_its_parent_is_blocked(env: dict[str, Any]) -> None:
    build_ok()
    path = env["data"] / "build" / "US" / "R985" / "R985-PARENT.json"
    parent = json.loads(path.read_text())
    parent["issues"].append({"code": "SchemaViolation", "message": "x", "severity": "ERROR"})
    path.write_text(json.dumps(parent))
    result = run("amazon", "submit", "--confirm", "R985-7", input="y\n")
    assert result.exit_code == 1
    assert env["fake"].listing_calls == []


def test_a_child_named_alone_is_held_until_its_parent_is_created(env: dict[str, Any]) -> None:
    build_ok()
    result = run("amazon", "submit", "--confirm", "R985-7", input="y\n")
    assert result.exit_code == 1
    assert env["fake"].listing_calls == []
    assert "R985-7" in result.output

    assert run("amazon", "submit", "--confirm", "R985-PARENT", input="y\n").exit_code == 0
    result = run("amazon", "submit", "--confirm", "R985-7", input="y\n")
    assert result.exit_code == 0, result.output
    assert [c.sku for c in env["fake"].writes()] == ["R985-PARENT", "R985-7"]


def test_a_child_named_alone_goes_when_its_live_parent_has_changed(env: dict[str, Any]) -> None:
    build_ok()
    assert run("amazon", "submit", "--confirm", "R985", input="y\n").exit_code == 0
    before = len(env["fake"].writes())
    for sku in ("R985-PARENT", "R985-7"):
        path = env["data"] / "build" / "US" / "R985" / f"{sku}.json"
        artifact = json.loads(path.read_text())
        artifact["attributes"]["item_name"][0]["value"] = "Edited After Review"
        path.write_text(json.dumps(artifact))

    result = run("amazon", "submit", "--confirm", "R985-7", input="y\n")
    assert result.exit_code == 0, result.output
    assert [c.sku for c in env["fake"].writes()[before:]] == ["R985-7"]


def test_feed_preview_sample_below_one_is_refused(env: dict[str, Any]) -> None:
    build_ok()
    result = run("amazon", "submit", "--confirm", "--feed", "--preview-sample", "0", input="y\n")
    assert result.exit_code == 2
    assert env["fake"].listing_calls == [] and env["fake"].feeds == {}


def test_a_missing_workbook_refuses_instead_of_selecting_everything(env: dict[str, Any]) -> None:
    build_ok()
    (env["root"] / "Product Listing.xlsx").unlink()
    for args in (("validate",), ("submit", "--confirm")):
        result = run("amazon", *args, input="y\n")
        assert result.exit_code == 2, result.output
        assert "name the SKUs explicitly" in result.output
    assert env["fake"].listing_calls == []


def test_a_failed_delete_keeps_the_listing_live(env: dict[str, Any]) -> None:
    build_ok()
    assert run("amazon", "submit", "--confirm", "E1154", input="y\n").exit_code == 0
    env["fake"].reject_delete.add("E1154")
    assert run("amazon", "delete", "E1154", "--confirm", input="y\n").exit_code == 1
    with Ledger(settings().state_db) as ledger:
        assert "E1154" in {e.sku for e in ledger.live_skus()}
    again = run("amazon", "submit", "--confirm", "E1154", input="y\n")
    assert again.exit_code == 0
    assert "up to date" in again.output


def test_delete_then_resubmit_sends_the_same_payload_again(env: dict[str, Any]) -> None:
    build_ok()
    assert run("amazon", "submit", "--confirm", "E1154", input="y\n").exit_code == 0
    writes = len(env["fake"].writes())
    assert run("amazon", "delete", "E1154", "--confirm", input="y\n").exit_code == 0
    assert env["fake"].deleted == ["E1154"]
    with Ledger(settings().state_db) as ledger:
        assert "E1154" not in {e.sku for e in ledger.live_skus()}
    assert run("amazon", "submit", "--confirm", "E1154", input="y\n").exit_code == 0
    assert len(env["fake"].writes()) == writes + 1
